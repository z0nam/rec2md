"""docx 회의록 렌더러 (extra `office`: python-docx, docxtpl).

- 양식에 원본 `form.docx`가 있으면 docxtpl(Jinja2 문법)로 채운다. 쓸 수 있는 변수:
  `doc_title`, 값 키마다 한 덩어리 글(`{{ place }}`, 여러 줄은 줄바꿈 유지),
  `lines.<값 키>`(줄 목록, `{% for x in lines.agenda %}`), `content`(회의내용 줄:
  `x.level`, `x.prefix`, `x.text`), `m`(Minutes 원본 dict).
- 원본이 없으면 python-docx로 `[[fields]]` 배치 그대로 표를 만든다(제목 + 라벨|값 표).
  같은 row의 필드는 한 줄에 나란히, 회의내용은 단계마다 들여쓴 문단.

pdf는 LibreOffice(`soffice --headless --convert-to pdf`), png는 그 pdf를 쪽마다 바꾼다.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from dataclasses import asdict
from pathlib import Path

from ..minutes import Minutes

FONT = "맑은 고딕"
FONT_PT = 10
LABEL_FILL = "EDEDED"
LEVEL_INDENT_MM = 4


def _need(mod: str):
    import importlib

    try:
        return importlib.import_module(mod)
    except ModuleNotFoundError:
        raise ModuleNotFoundError(
            f"docx 렌더에 {mod}가 필요함: uv sync --extra office (pip install 'rec2md[office]')") from None


def render_docx(m: Minutes, template, out: Path) -> Path:
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    form = template.file("docx")
    if form is not None:
        return _fill_docxtpl(m, template, form, out)
    return _build(m, template, out)


# ── 사용자 양식(form.docx) ─────────────────────────────────────────────

def context(m: Minutes, template) -> dict:
    """docxtpl에 넘기는 변수. 값 계산은 Template에 맡긴다(형식 간 표기 통일)."""
    vals = template.values(m)
    ctx = {k: "\n".join(v) for k, v in vals.items()}
    ctx.update(doc_title=template.doc_title(m), lines=vals, m=asdict(m),
               content=[{"level": x.level, "prefix": x.prefix, "text": x.text, "ts": x.ts}
                        for x in template.content_lines(m)])
    return ctx


def _fill_docxtpl(m: Minutes, template, form: Path, out: Path) -> Path:
    docxtpl = _need("docxtpl")
    doc = docxtpl.DocxTemplate(str(form))
    ctx = context(m, template)
    # 여러 줄 값은 Listing으로 감싸야 줄바꿈이 살아난다
    ctx.update({k: docxtpl.Listing(v) for k, v in ctx.items() if isinstance(v, str) and "\n" in v})
    doc.render(ctx, autoescape=True)
    doc.save(str(out))
    return out


# ── 원본 없이 만들기 ──────────────────────────────────────────────────

def _build(m: Minutes, template, out: Path) -> Path:
    docx = _need("docx")
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt

    doc = docx.Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin = sec.bottom_margin = Mm(20)
    _set_font(doc.styles["Normal"], FONT_PT)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(template.doc_title(m))
    run.bold, run.font.size = True, Pt(16)

    vals = template.values(m)
    rows = template.rows()
    pairs = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=2 * pairs)
    table.style = "Table Grid"
    text_w = sec.page_width - sec.left_margin - sec.right_margin
    label_w = Mm(24)
    value_w = int((text_w - pairs * label_w) / pairs)
    table.autofit = False
    for ci, col in enumerate(table.columns):  # LibreOffice·Word는 칸 폭보다 열 폭(gridCol)을 먼저 본다
        col.width = label_w if ci % 2 == 0 else value_w
    for ri, row in enumerate(rows):
        cells = table.rows[ri].cells
        for fi, f in enumerate(row):
            label, value = cells[2 * fi], cells[2 * fi + 1]
            label.width, value.width = label_w, value_w
            if fi == len(row) - 1 and 2 * fi + 1 < 2 * pairs - 1:  # 마지막 값 칸이 남은 폭을 차지
                value = value.merge(cells[2 * pairs - 1])
                value.width = text_w - fi * (label_w + value_w) - label_w
            _shade(label, LABEL_FILL)
            label.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = label.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(f.label).bold = True
            if f.value == "content":
                lines = [(x.level, str(x), x.prefix) for x in template.content_lines(m)]
                value.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            else:
                lines = [(1, s, "- ") if f.value == "attendees" and s.startswith("- ") else (0, s)
                         for s in vals[f.value]]
                value.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _fill_cell(value, lines)
    doc.core_properties.title = m.title
    doc.core_properties.author = ""
    doc.save(str(out))
    return out


def _fill_cell(cell, lines: list[tuple]) -> None:
    """줄마다 문단 하나. (단계, 글[, 머리기호]) — 머리기호가 있으면 둘째 줄이 기호 뒤 글자에 맞게 내어쓴다."""
    from docx.shared import Mm, Pt

    first = cell.paragraphs[0]
    for i, (level, text, *prefix) in enumerate(lines or [(0, "")]):
        p = first if i == 0 else cell.add_paragraph()
        p.add_run(text)
        left = Mm(LEVEL_INDENT_MM * level)
        hang = Pt(_marker_width_pt(prefix[0], FONT_PT)) if prefix and prefix[0].strip() else 0
        if left or hang:
            p.paragraph_format.left_indent = left + hang
            p.paragraph_format.first_line_indent = -hang if hang else None


def _marker_width_pt(prefix: str, size: float) -> float:
    """머리기호 폭(pt) 어림. 도형·화살표·한글은 글자 크기만큼, 숫자·기호·공백은 절반 (hwpx와 같은 규칙)."""
    return sum(size if ord(ch) >= 0x2190 else size * 0.5 for ch in prefix)


def _set_font(style, size_pt: int) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    style.font.name = FONT
    style.font.size = Pt(size_pt)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)  # 한글 글꼴은 eastAsia로 따로 지정해야 먹는다


def _shade(cell, fill: str) -> None:
    from docx.oxml.ns import qn

    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): fill})
    tcpr.append(shd)


# ── pdf·png ────────────────────────────────────────────────────────────

SOFFICE_CANDIDATES = ["soffice", "libreoffice", "/opt/homebrew/bin/soffice",
                      "/Applications/LibreOffice.app/Contents/MacOS/soffice",
                      r"C:\Program Files\LibreOffice\program\soffice.exe"]


def find_soffice() -> str | None:
    if env := os.environ.get("REC2MD_SOFFICE"):
        return env
    for c in SOFFICE_CANDIDATES:
        if os.path.isabs(c) and os.path.exists(c):
            return c
        if found := shutil.which(c):
            return found
    return None


def to_pdf(src: Path, out: Path) -> Path:
    from ..export import ExportError

    soffice = find_soffice()
    if soffice is None:
        raise ExportError("docx → pdf에 LibreOffice가 필요함: brew install --cask libreoffice / "
                          "apt install libreoffice-writer, 또는 REC2MD_SOFFICE로 경로 지정")
    with tempfile.TemporaryDirectory(prefix="rec2md-soffice-") as td:
        # 사용자 프로필을 따로 써야 이미 떠 있는 LibreOffice와 부딪치지 않는다
        r = subprocess.run([soffice, f"-env:UserInstallation={Path(td, 'profile').as_uri()}", "--headless",
                            "--convert-to", "pdf", "--outdir", td, str(src)],
                           capture_output=True, text=True, timeout=300)
        made = Path(td) / (src.stem + ".pdf")
        if r.returncode != 0 or not made.is_file():
            raise ExportError(f"LibreOffice pdf 변환 실패: {(r.stderr or r.stdout).strip()[-500:]}")
        out.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(made, out)
    return out


def to_png(src: Path, out_dir: Path) -> list[Path]:
    from ..export import pdf_to_png

    stem = src.name[: -len(src.suffix)]
    with tempfile.TemporaryDirectory(prefix="rec2md-docx-") as td:
        pdf = to_pdf(src, Path(td) / f"{stem}.pdf")
        out_dir.mkdir(parents=True, exist_ok=True)
        return pdf_to_png(pdf, out_dir, stem)
